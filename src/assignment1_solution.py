"""Reproducible Assignment 1 data cleaning and Gaussian Naive Bayes solution."""

from __future__ import annotations

import argparse
import csv
import json
import math
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook


MISSING_TOKENS = {"", "?", "na", "n/a", "nan", "null", "missing"}
FEATURE_NAMES = ("A", "B", "C", "D")


def read_rows(path: Path) -> list[tuple[Any, ...]]:
    """Read the first worksheet without changing its stored values."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        return [tuple(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def is_missing(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip().lower() in MISSING_TOKENS
    if isinstance(value, (int, float, np.number)):
        return not math.isfinite(float(value))
    return False


def as_number(value: Any, location: str) -> float:
    if is_missing(value):
        return math.nan
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"Nonnumeric value at {location}: {value!r}") from error
    if not math.isfinite(number):
        raise ValueError(f"Nonfinite value at {location}: {value!r}")
    return number


def class_counts(labels: np.ndarray) -> dict[str, int]:
    values, counts = np.unique(labels.astype(int), return_counts=True)
    return {str(int(value)): int(count) for value, count in zip(values, counts)}


def run_analysis(source_root: Path) -> dict[str, Any]:
    assignment = source_root / "assigenment1"
    train_rows = read_rows(assignment / "TrainingData (2).xlsx")
    test_rows = read_rows(assignment / "TestData.xlsx")

    if len(train_rows) != 120 or any(len(row) != 5 for row in train_rows):
        raise ValueError("TrainingData (2).xlsx must contain 120 rows and 5 columns.")
    if len(test_rows) != 30 or any(len(row) != 4 for row in test_rows):
        raise ValueError("TestData.xlsx must contain 30 rows and 4 columns.")

    missing_cells: list[dict[str, Any]] = []
    x_train = np.empty((len(train_rows), 4), dtype=float)
    y_train = np.empty(len(train_rows), dtype=int)
    for row_index, row in enumerate(train_rows, start=1):
        for column_index, value in enumerate(row[:4]):
            if is_missing(value):
                missing_cells.append({"row": row_index, "feature": FEATURE_NAMES[column_index]})
                x_train[row_index - 1, column_index] = math.nan
            else:
                x_train[row_index - 1, column_index] = as_number(
                    value, f"training row {row_index}, feature {FEATURE_NAMES[column_index]}"
                )
        if is_missing(row[4]):
            missing_cells.append({"row": row_index, "feature": "class label E"})
            raise ValueError(f"Missing class label at training row {row_index}.")
        label_value = as_number(row[4], f"training row {row_index}, class label E")
        if not label_value.is_integer():
            raise ValueError(f"Noninteger class label at training row {row_index}.")
        y_train[row_index - 1] = int(label_value)

    x_test = np.array(
        [
            [as_number(value, f"test row {i}, feature {FEATURE_NAMES[j]}") for j, value in enumerate(row)]
            for i, row in enumerate(test_rows, start=1)
        ],
        dtype=float,
    )

    q1 = np.empty(4, dtype=float)
    q3 = np.empty(4, dtype=float)
    lower = np.empty(4, dtype=float)
    upper = np.empty(4, dtype=float)
    median = np.empty(4, dtype=float)
    mad = np.empty(4, dtype=float)
    iqr_rows: set[int] = set()
    mad_rows: set[int] = set()
    outlier_details: list[dict[str, Any]] = []

    for feature_index, feature_name in enumerate(FEATURE_NAMES):
        column = x_train[:, feature_index]
        observed = column[np.isfinite(column)]
        if observed.size < 4:
            raise ValueError(f"Too few observed values in feature {feature_name}.")
        q1[feature_index], q3[feature_index] = np.quantile(
            observed, (0.25, 0.75), method="linear"
        )
        spread = q3[feature_index] - q1[feature_index]
        lower[feature_index] = q1[feature_index] - 1.5 * spread
        upper[feature_index] = q3[feature_index] + 1.5 * spread
        median[feature_index] = np.median(observed)
        mad[feature_index] = np.median(np.abs(observed - median[feature_index]))

        for zero_based_row, value in enumerate(column):
            if not math.isfinite(float(value)):
                continue
            if value < lower[feature_index] or value > upper[feature_index]:
                row_number = zero_based_row + 1
                iqr_rows.add(row_number)
                outlier_details.append(
                    {"row": row_number, "feature": feature_name, "value": float(value)}
                )
            if mad[feature_index] > 0:
                robust_z = 0.6744897501960817 * abs(value - median[feature_index]) / mad[feature_index]
                if robust_z > 3.5:
                    mad_rows.add(zero_based_row + 1)

    classwise_iqr_rows: set[int] = set()
    classwise_iqr_details: list[dict[str, Any]] = []
    for class_value in np.unique(y_train):
        class_mask = y_train == class_value
        for feature_index, feature_name in enumerate(FEATURE_NAMES):
            column = x_train[class_mask, feature_index]
            observed = column[np.isfinite(column)]
            class_q1, class_q3 = np.quantile(observed, (0.25, 0.75), method="linear")
            class_fence = 1.5 * (class_q3 - class_q1)
            for zero_based_row in np.flatnonzero(class_mask):
                value = x_train[zero_based_row, feature_index]
                if math.isfinite(float(value)) and (
                    value < class_q1 - class_fence or value > class_q3 + class_fence
                ):
                    row_number = int(zero_based_row + 1)
                    classwise_iqr_rows.add(row_number)
                    classwise_iqr_details.append(
                        {"row": row_number, "class": int(class_value), "feature": feature_name, "value": float(value)}
                    )

    missing_rows = {int(cell["row"]) for cell in missing_cells}
    deleted_rows = sorted(missing_rows | iqr_rows)
    retained_mask = np.ones(len(train_rows), dtype=bool)
    if deleted_rows:
        retained_mask[np.asarray(deleted_rows, dtype=int) - 1] = False
    x_clean = x_train[retained_mask]
    y_clean = y_train[retained_mask]
    if not np.isfinite(x_clean).all():
        raise ValueError("Cleaning left a missing feature value in the training matrix.")

    classes = np.unique(y_clean)
    priors: list[float] = []
    means: list[list[float]] = []
    variances: list[list[float]] = []
    for class_value in classes:
        class_features = x_clean[y_clean == class_value]
        priors.append(float(class_features.shape[0] / x_clean.shape[0]))
        means.append(class_features.mean(axis=0).tolist())
        class_variance = class_features.var(axis=0, ddof=0)
        if np.any(class_variance <= 0):
            raise ValueError(f"Zero or negative class-conditional variance for class {class_value}.")
        variances.append(class_variance.tolist())

    means_array = np.asarray(means, dtype=float)
    variances_array = np.asarray(variances, dtype=float)
    priors_array = np.asarray(priors, dtype=float)
    log_scores = np.empty((len(x_test), len(classes)), dtype=float)
    for class_index in range(len(classes)):
        log_scores[:, class_index] = math.log(priors_array[class_index]) - 0.5 * np.sum(
            np.log(2.0 * math.pi * variances_array[class_index])
            + (x_test - means_array[class_index]) ** 2 / variances_array[class_index],
            axis=1,
        )
    predictions = classes[np.argmax(log_scores, axis=1)]
    sorted_scores = np.sort(log_scores, axis=1)
    score_margins = sorted_scores[:, -1] - sorted_scores[:, -2]

    thresholds = [
        {
            "feature": FEATURE_NAMES[j],
            "q1": float(q1[j]),
            "q3": float(q3[j]),
            "iqr": float(q3[j] - q1[j]),
            "lower": float(lower[j]),
            "upper": float(upper[j]),
            "median": float(median[j]),
            "mad": float(mad[j]),
        }
        for j in range(4)
    ]
    return {
        "source_root": str(source_root),
        "training_shape": list(x_train.shape),
        "test_shape": list(x_test.shape),
        "feature_names": list(FEATURE_NAMES),
        "missing_cells": missing_cells,
        "missing_rows": sorted(missing_rows),
        "outlier_rule": "Per-feature global Tukey IQR; NumPy linear quantiles; 1.5*IQR fences.",
        "quantile_method": "linear interpolation, method='linear'",
        "thresholds": thresholds,
        "iqr_outlier_rows": sorted(iqr_rows),
        "classwise_iqr_outlier_rows_sensitivity_only": sorted(classwise_iqr_rows),
        "classwise_iqr_details_sensitivity_only": sorted(
            classwise_iqr_details, key=lambda item: (item["row"], item["feature"])
        ),
        "mad_robust_z_threshold": 3.5,
        "mad_outlier_rows": sorted(mad_rows),
        "outlier_details": sorted(outlier_details, key=lambda item: (item["row"], item["feature"])),
        "deleted_rows": deleted_rows,
        "deleted_count": len(deleted_rows),
        "retained_count": int(x_clean.shape[0]),
        "class_counts_before": class_counts(y_train),
        "class_counts_after": class_counts(y_clean),
        "model": {
            "name": "Gaussian Naive Bayes",
            "variance_estimator": "Maximum-likelihood population variance (ddof=0; divisor n_k).",
            "classes": classes.astype(int).tolist(),
            "class_counts": class_counts(y_clean),
            "priors": priors,
            "means": means,
            "variances": variances,
            "predictions": predictions.astype(int).tolist(),
            "log_discriminants": log_scores.tolist(),
            "minimum_top_two_log_score_margin": float(score_margins.min()),
            "test_count": int(len(predictions)),
        },
    }


def write_outputs(source_root: Path, output_dir: Path) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    result = run_analysis(source_root)
    json_path = output_dir / "assignment1_results.json"
    json_path.write_text(json.dumps(result, indent=2), encoding="utf-8")
    csv_path = output_dir / "assignment1_predictions.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(("test_record", "predicted_class", "log_score_class_1", "log_score_class_2", "log_score_class_3"))
        scores = result["model"]["log_discriminants"]
        for index, (label, row_scores) in enumerate(zip(result["model"]["predictions"], scores), start=1):
            writer.writerow((index, label, *[f"{score:.12g}" for score in row_scores]))
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    result = write_outputs(args.source_root, args.output_dir)
    print(json.dumps({
        "missing_rows": result["missing_rows"],
        "iqr_outlier_rows": result["iqr_outlier_rows"],
        "mad_outlier_rows": result["mad_outlier_rows"],
        "deleted_rows": result["deleted_rows"],
        "retained_count": result["retained_count"],
        "class_counts_after": result["class_counts_after"],
        "predictions": result["model"]["predictions"],
    }, indent=2))


if __name__ == "__main__":
    main()
